"""Remember what the accountant decided so next month's run does not ask again.

Decisions are stored per supplier and keyed on the normalised invoice numbers on
both sides, so a pair confirmed in October still applies when the same documents
show up again in a later or revised run.

The usual loop is:

1. ``recon2b run`` writes a report whose Review sheet has a Decision column.
2. The accountant types Accept or Reject against the pairs they checked.
3. ``recon2b learn report.xlsx`` reads those answers into ``decisions.json``.
4. The next run applies them before any automatic matching.
"""

from __future__ import annotations

import datetime as dt
import json
from pathlib import Path

import polars as pl
from openpyxl import load_workbook

from .matcher import Decisions
from .normalise import invoice_keys

ACCEPT, REJECT = "accept", "reject"
REVIEW_SHEET = "Review"


def _key(d: dict) -> tuple[str, str, str, str]:
    return (d["gstin"], d["doc_type"], d["books_inv"], d["gstr2b_inv"])


class DecisionStore:
    def __init__(self, path: str | Path | None = None):
        self.path = Path(path) if path else None
        self.items: dict[tuple[str, str, str, str], dict] = {}
        if self.path and self.path.exists():
            data = json.loads(self.path.read_text() or "{}")
            for d in data.get("decisions", []):
                self.items[_key(d)] = d

    def __len__(self) -> int:
        return len(self.items)

    def add(
        self,
        gstin: str,
        doc_type: str,
        books_inv: str,
        gstr2b_inv: str,
        decision: str,
        normalise: bool = True,
    ) -> None:
        decision = decision.strip().lower()
        if decision not in (ACCEPT, REJECT):
            raise ValueError(f"decision must be accept or reject, got {decision!r}")
        if normalise:
            books_inv, gstr2b_inv = _norm(books_inv), _norm(gstr2b_inv)
        d = {
            "gstin": gstin.strip().upper(),
            "doc_type": doc_type.strip().upper(),
            "books_inv": books_inv,
            "gstr2b_inv": gstr2b_inv,
            "decision": decision,
            "decided_on": dt.date.today().isoformat(),
        }
        self.items[_key(d)] = d

    def save(self, path: str | Path | None = None) -> Path:
        target = Path(path) if path else self.path
        if target is None:
            raise ValueError("no path to save decisions to")
        body = {"version": 1, "decisions": sorted(self.items.values(), key=_key)}
        target.write_text(json.dumps(body, indent=2))
        return target

    def to_matcher(self) -> Decisions:
        def frame(kind: str) -> pl.DataFrame:
            rows = [d for d in self.items.values() if d["decision"] == kind]
            return pl.DataFrame(
                {
                    "gstin": [d["gstin"] for d in rows],
                    "doc_type": [d["doc_type"] for d in rows],
                    "b_norm": [d["books_inv"] for d in rows],
                    "t_norm": [d["gstr2b_inv"] for d in rows],
                },
                schema={
                    "gstin": pl.Utf8,
                    "doc_type": pl.Utf8,
                    "b_norm": pl.Utf8,
                    "t_norm": pl.Utf8,
                },
            )

        return Decisions(accepted=frame(ACCEPT), rejected=frame(REJECT))

    def learn_from_report(self, report: str | Path) -> int:
        """Read the Decision column of a report's Review sheet. Returns how many were read."""
        wb = load_workbook(report, read_only=True, data_only=True)
        try:
            if REVIEW_SHEET not in wb.sheetnames:
                raise ValueError(f"{Path(report).name} has no '{REVIEW_SHEET}' sheet")
            rows = wb[REVIEW_SHEET].iter_rows(values_only=True)
            header = [str(h or "").strip() for h in next(rows)]
            col = {name: i for i, name in enumerate(header)}
            needed = ["Decision", "gstin", "doc_type", "books_inv_norm", "gstr2b_inv_norm"]
            missing = [n for n in needed if n not in col]
            if missing:
                raise ValueError(f"Review sheet is missing columns: {', '.join(missing)}")
            count = 0
            for r in rows:
                choice = str(r[col["Decision"]] or "").strip().lower()
                if choice not in (ACCEPT, REJECT):
                    continue
                self.add(
                    str(r[col["gstin"]]),
                    str(r[col["doc_type"]]),
                    str(r[col["books_inv_norm"]]),
                    str(r[col["gstr2b_inv_norm"]]),
                    choice,
                    normalise=False,
                )
                count += 1
            return count
        finally:
            wb.close()


def _norm(inv: str) -> str:
    # GST began in 2017, so any financial year since then may be glued into a number
    years = list(range(2017, dt.date.today().year + 1))
    out = pl.DataFrame({"inv_no": [inv]}).select(invoice_keys("inv_no", years))
    return out["inv_norm"][0]
