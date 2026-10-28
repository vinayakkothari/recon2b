import datetime as dt
import json

import pytest
from openpyxl import Workbook

from conftest import SUP_A, SUP_B
from recon2b.loaders import load_books, load_gstr2b
from recon2b.loaders.tabular import map_columns


def test_books_csv_with_tally_style_headers(tmp_path):
    p = tmp_path / "books.csv"
    p.write_text(
        "Date,Particulars,GSTIN/UIN,Supplier Invoice No.,Voucher Type,Taxable Value,"
        "IGST,CGST,SGST/UTGST,Cess,Gross Total\n"
        f'05-10-2026,Acme,{SUP_A},INV/26-27/1,Purchase,"1,000.00",0,90,90,0,"1,180.00"\n'
        f"06-10-2026,Acme,{SUP_A.lower()},CN/1,Credit Note,500,0,45,45,0,590\n"
        f"07-10-2026,Beta,{SUP_B},77,Purchase,(200),36,0,0,0,236\n"
    )
    df = load_books(p)
    assert df["gstin"].to_list() == [SUP_A, SUP_A, SUP_B]  # case and spacing cleaned
    assert df["inv_no"].to_list() == ["INV/26-27/1", "CN/1", "77"]
    assert df["doc_type"].to_list() == ["INV", "CRN", "CRN"]  # negative amount -> credit note
    assert df["taxable"].to_list() == [1000.0, 500.0, 200.0]
    assert df["cgst"].to_list() == [90.0, 45.0, 0.0]
    assert df["inv_date"][0] == dt.date(2026, 10, 5)


def test_column_map_override(tmp_path):
    p = tmp_path / "odd.csv"
    p.write_text(f"Vendor Tax Id,Bill Ref,Bill Dt,Net\n{SUP_A},X1,01-10-2026,100\n")
    with pytest.raises(ValueError, match="could not find columns"):
        load_books(p)
    df = load_books(
        p,
        column_map={
            "gstin": "Vendor Tax Id",
            "inv_no": "Bill Ref",
            "inv_date": "Bill Dt",
            "taxable": "Net",
        },
    )
    assert df.row(0, named=True)["inv_no"] == "X1"


def test_alias_priority_prefers_supplier_invoice_date():
    m = map_columns(["Date", "Supplier Invoice Date", "GSTIN", "Invoice No", "Taxable Value"])
    assert m["inv_date"] == "Supplier Invoice Date"


def test_books_excel_with_title_rows(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Purchase Register"
    ws.append(["Acme Industries Pvt Ltd"])
    ws.append(["Purchase Register 1-Oct-2026 to 31-Oct-2026"])
    ws.append([])
    ws.append(["Date", "Party Name", "GSTIN/UIN", "Supplier Invoice No", "Taxable Value", "IGST"])
    ws.append([dt.datetime(2026, 10, 3), "Acme", SUP_A, 101.0, 2500, 450])
    ws.append([None, None, None, None, None, None])
    ws.append([dt.datetime(2026, 10, 4), "Acme", SUP_A, "A-2", 100.5, 18.09])
    path = tmp_path / "books.xlsx"
    wb.save(path)
    df = load_books(path)
    assert df.height == 2
    assert df["inv_no"].to_list() == ["101", "A-2"]  # 101.0 from Excel is not "101.0"
    assert df["inv_date"].to_list() == [dt.date(2026, 10, 3), dt.date(2026, 10, 4)]


def _portal_json():
    return {
        "chksum": "x",
        "data": {
            "gstin": "29AAAAA0000A1Z5",
            "rtnprd": "102026",
            "docdata": {
                "b2b": [
                    {
                        "ctin": SUP_A,
                        "trdnm": "Acme",
                        "supfildt": "11-11-2026",
                        "supprd": "102026",
                        "inv": [
                            {
                                "inum": "INV/26-27/1",
                                "typ": "R",
                                "dt": "05-10-2026",
                                "val": 1180,
                                "pos": "29",
                                "rev": "N",
                                "itcavl": "Y",
                                "rsn": "",
                                "diffprcnt": 1,
                                "items": [
                                    {
                                        "num": 1,
                                        "rt": 18,
                                        "txval": 600,
                                        "igst": 0,
                                        "cgst": 54,
                                        "sgst": 54,
                                        "cess": 0,
                                    },
                                    {
                                        "num": 2,
                                        "rt": 18,
                                        "txval": 400,
                                        "igst": 0,
                                        "cgst": 36,
                                        "sgst": 36,
                                        "cess": 0,
                                    },
                                ],
                            },
                            {
                                "inum": "INV/26-27/2",
                                "dt": "06-10-2026",
                                "val": 118,
                                "rev": "Y",
                                "itcavl": "N",
                                "txval": 100,
                                "cgst": 9,
                                "sgst": 9,
                            },
                        ],
                    }
                ],
                "cdnr": [
                    {
                        "ctin": SUP_B,
                        "trdnm": "Beta",
                        "nt": [
                            {
                                "ntnum": "CN-4",
                                "typ": "C",
                                "dt": "07-10-2026",
                                "val": 59,
                                "items": [{"txval": 50, "igst": 9}],
                            },
                            {
                                "ntnum": "DN-1",
                                "typ": "D",
                                "dt": "08-10-2026",
                                "val": 11.8,
                                "items": [{"txval": 10, "igst": 1.8}],
                            },
                        ],
                    }
                ],
            },
        },
    }


def test_gstr2b_json(tmp_path):
    p = tmp_path / "2b.json"
    p.write_text(json.dumps(_portal_json()))
    df = load_gstr2b(p)
    assert df["inv_no"].to_list() == ["INV/26-27/1", "INV/26-27/2", "CN-4", "DN-1"]
    assert df["doc_type"].to_list() == ["INV", "INV", "CRN", "DBN"]
    assert df["taxable"].to_list() == [1000.0, 100.0, 50.0, 10.0]  # items summed
    assert df["cgst"].to_list()[:2] == [90.0, 9.0]  # totals on the invoice when no items
    assert df["itc_available"].to_list() == [True, False, True, True]
    assert df["reverse_charge"].to_list() == [False, True, False, False]
    assert df["inv_date"][0] == dt.date(2026, 10, 5)


def test_several_gstr2b_files_get_unique_rows(tmp_path):
    a, b = tmp_path / "a.json", tmp_path / "b.json"
    a.write_text(json.dumps(_portal_json()))
    b.write_text(json.dumps(_portal_json()))
    df = load_gstr2b([a, b])
    assert df.height == 8 and df["row"].n_unique() == 8


def test_gstr2b_portal_excel_layout(tmp_path):
    """Mimics the portal's download: title rows, two-row headers, extra sheets."""
    wb = Workbook()
    wb.active.title = "Read me"
    wb.active.append(["Goods and Services Tax - GSTR 2B"])
    wb.create_sheet("ITC Available").append(["Summary", "Integrated Tax"])
    ws = wb.create_sheet("B2B")
    ws.append(["Goods and Services Tax - GSTR-2B"])
    ws.append(["Taxable inward supplies received from registered persons"])
    ws.append([])
    ws.append(
        [
            "GSTIN of supplier",
            "Trade/Legal name",
            "Invoice details",
            None,
            None,
            None,
            "Place of supply",
            "Supply Attract Reverse Charge",
            "Taxable Value (₹)",
            "Tax Amount",
            None,
            None,
            None,
            "GSTR-1/IFF/GSTR-5 Period",
            "GSTR-1/IFF/GSTR-5 Filing Date",
            "ITC Availability",
            "Reason",
        ]
    )
    ws.append(
        [
            None,
            None,
            "Invoice number",
            "Invoice type",
            "Invoice Date",
            "Invoice Value(₹)",
            None,
            None,
            None,
            "Integrated Tax(₹)",
            "Central Tax(₹)",
            "State/UT Tax(₹)",
            "Cess(₹)",
        ]
    )
    ws.append(
        [
            SUP_A,
            "Acme",
            "INV/26-27/1",
            "Regular",
            "05-10-2026",
            1180,
            "29-Karnataka",
            "No",
            1000,
            0,
            90,
            90,
            0,
            "Oct'26",
            "11-11-2026",
            "Yes",
            None,
        ]
    )
    ws.append(
        [
            SUP_A,
            "Acme",
            "INV/26-27/2",
            "Regular",
            "06-10-2026",
            118,
            "29-Karnataka",
            "Yes",
            100,
            0,
            9,
            9,
            0,
            "Oct'26",
            "11-11-2026",
            "No",
            "POS",
        ]
    )
    cd = wb.create_sheet("B2B-CDNR")
    cd.append(
        [
            "GSTIN of supplier",
            "Trade/Legal name",
            "Credit note/Debit note details",
            None,
            None,
            None,
            None,
            "Place of supply",
            "Supply Attract Reverse Charge",
            "Taxable Value (₹)",
            "Tax Amount",
        ]
    )
    cd.append(
        [
            None,
            None,
            "Note number",
            "Note type",
            "Note Supply type",
            "Note date",
            "Note Value (₹)",
            None,
            None,
            None,
            "Integrated Tax(₹)",
            "Central Tax(₹)",
            "State/UT Tax(₹)",
            "Cess(₹)",
        ]
    )
    cd.append(
        [
            SUP_B,
            "Beta",
            "CN-4",
            "Credit Note",
            "Regular",
            "07-10-2026",
            59,
            "29",
            "No",
            50,
            9,
            0,
            0,
            0,
        ]
    )
    path = tmp_path / "2b.xlsx"
    wb.save(path)

    df = load_gstr2b(path)
    assert df["inv_no"].to_list() == ["INV/26-27/1", "INV/26-27/2", "CN-4"]
    assert df["gstin"].to_list() == [SUP_A, SUP_A, SUP_B]
    assert df["doc_type"].to_list() == ["INV", "INV", "CRN"]
    assert df["taxable"].to_list() == [1000.0, 100.0, 50.0]
    assert df["igst"].to_list() == [0.0, 0.0, 9.0]
    assert df["itc_available"].to_list() == [True, False, True]
    assert df["reverse_charge"].to_list() == [False, True, False]


def test_unsupported_file(tmp_path):
    p = tmp_path / "x.pdf"
    p.write_text("")
    with pytest.raises(ValueError, match="unsupported"):
        load_gstr2b(p)
    with pytest.raises(ValueError, match="unsupported"):
        load_books(p)
