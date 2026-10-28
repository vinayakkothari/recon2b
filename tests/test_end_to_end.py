"""Whole flows: synthetic data in, report out, decisions learned, CLI and web."""

import polars as pl
import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from recon2b import synth
from recon2b.cli import app
from recon2b.decisions import DecisionStore
from recon2b.loaders import load_books, load_gstr2b
from recon2b.matcher import reconcile
from recon2b.pipeline import run


@pytest.fixture(scope="module")
def data(tmp_path_factory):
    return synth.write(tmp_path_factory.mktemp("synth"), 5000, seed=3)


def test_accuracy_on_synthetic_data(data):
    res = reconcile(load_books(data["books"]), load_gstr2b(data["gstr2b"]))
    score = synth.evaluate(
        res.rows, pl.read_csv(data["truth_books"]), pl.read_csv(data["truth_2b"])
    )
    assert score["precision"] >= 0.995
    assert score["recall"] >= 0.99


def test_every_input_row_appears_exactly_once(data):
    b, t = load_books(data["books"]), load_gstr2b(data["gstr2b"])
    rows = reconcile(b, t).rows
    assert sorted(rows["books_row"].drop_nulls().to_list()) == list(range(b.height))
    assert sorted(rows["gstr2b_row"].drop_nulls().to_list()) == list(range(t.height))


def test_report_and_learning_loop(data, tmp_path):
    report = tmp_path / "report.xlsx"
    result, _ = run(data["books"], [data["gstr2b"]], report)
    wb = load_workbook(report)
    assert wb.sheetnames[:4] == ["Summary", "Reconciliation", "Review", "Suppliers"]
    review = wb["Review"]
    header = [c.value for c in review[1]]
    assert header[0] == "Decision"
    n_review = review.max_row - 1
    assert n_review == int(result.rows["needs_review"].sum()) > 0

    # the accountant accepts the first pair and rejects the second
    review.cell(row=2, column=1, value="Accept")
    review.cell(row=3, column=1, value="Reject")
    wb.save(report)

    decisions = tmp_path / "decisions.json"
    store = DecisionStore(decisions)
    assert store.learn_from_report(report) == 2
    store.save()

    again, _ = run(data["books"], [data["gstr2b"]], tmp_path / "again.xlsx", decisions=decisions)
    assert again.rows.filter(pl.col("match_pass") == "manual").height == 1
    assert int(again.rows["needs_review"].sum()) < n_review


def test_flat_output(data, tmp_path):
    result, path = run(data["books"], [data["gstr2b"]], tmp_path / "out.parquet")
    back = pl.read_parquet(path)
    assert back.height == result.rows.height and "status_label" in back.columns


def test_cli(tmp_path):
    runner = CliRunner()
    folder = tmp_path / "d"
    out = runner.invoke(app, ["synth", "--out-dir", str(folder), "--invoices", "500"])
    assert out.exit_code == 0, out.output
    report = tmp_path / "r.xlsx"
    out = runner.invoke(
        app,
        [
            "run",
            "-b",
            str(folder / "books.csv"),
            "-g",
            str(folder / "gstr2b.json"),
            "-o",
            str(report),
        ],
    )
    assert out.exit_code == 0, out.output
    assert "MISSING_IN_2B" in out.output and report.exists()
    out = runner.invoke(app, ["eval", str(folder)])
    assert out.exit_code == 0 and "precision" in out.output
    out = runner.invoke(
        app,
        [
            "pair",
            "--gstin",
            "29AAACA1111A1ZX",
            "--books-inv",
            "inv 1",
            "--gstr2b-inv",
            "INV/26-27/001",
            "--decisions",
            str(tmp_path / "d.json"),
        ],
    )
    assert out.exit_code == 0
    saved = DecisionStore(tmp_path / "d.json")
    (only,) = saved.items.values()
    assert only["books_inv"] == "INV1" and only["gstr2b_inv"] == "INV1"


def test_web_upload(data):
    fastapi = pytest.importorskip("fastapi")
    from fastapi.testclient import TestClient

    from recon2b.web import app as web_app

    assert fastapi
    client = TestClient(web_app)
    assert "recon2b" in client.get("/").text
    with open(data["books"], "rb") as b, open(data["gstr2b"], "rb") as g:
        res = client.post(
            "/reconcile",
            files=[
                ("books", ("books.csv", b, "text/csv")),
                ("gstr2b", ("gstr2b.json", g, "application/json")),
            ],
            data={"tolerance": "1"},
        )
    assert res.status_code == 200
    assert res.content[:2] == b"PK"  # an xlsx is a zip file

    bad = client.post(
        "/reconcile",
        files=[
            ("books", ("x.pdf", b"x", "application/pdf")),
            ("gstr2b", ("y.json", b"{}", "application/json")),
        ],
    )
    assert bad.status_code == 422
